
import io
import re
import hmac
import secrets
from datetime import datetime

import pandas as pd
import streamlit as st
import openpyxl

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN

from openpyxl.utils import column_index_from_string
from google import genai
from google.genai import types
from supabase import create_client


# =========================================================
# 1. CẤU HÌNH TRANG
# =========================================================

st.set_page_config(
    page_title="AI Smart Eval THCS",
    page_icon="📝",
    layout="centered"
)

st.markdown(
    """
    <style>
    h1 {font-size: 2.15rem !important;}
    @media (max-width: 600px) {
        h1 {font-size: 1.7rem !important;}
    }
    </style>
    """,
    unsafe_allow_html=True
)

st.title("📝 AI SMART EVAL THCS")
st.caption(
    "Hệ thống hỗ trợ chấm bài thực hành Excel, Word, PowerPoint "
    "theo barem giáo viên thiết lập và phản hồi sư phạm bằng AI"
)


# =========================================================
# 2. SECRETS + DATABASE
# =========================================================

def lay_secret(ten, mac_dinh=None):
    try:
        return st.secrets[ten]
    except Exception:
        return mac_dinh


TEACHER_PIN = str(lay_secret("TEACHER_PIN", ""))
GEMINI_API_KEY = str(lay_secret("GEMINI_API_KEY", ""))
GEMINI_MODEL = str(lay_secret("GEMINI_MODEL", "gemini-3.8-flash"))
SUPABASE_URL = str(lay_secret("SUPABASE_URL", ""))
SUPABASE_KEY = str(lay_secret("SUPABASE_SERVICE_KEY", ""))


@st.cache_resource
def ket_noi_supabase():
    if not SUPABASE_URL or not SUPABASE_KEY:
        return None
    return create_client(SUPABASE_URL, SUPABASE_KEY)


db = ket_noi_supabase()


# =========================================================
# 3. HÀM DATABASE
# =========================================================

def tao_ma_ngan():
    alphabet = "ABCDEFGHJKLMNPQRSTUVWXYZ23456789"
    return "AI-" + "".join(secrets.choice(alphabet) for _ in range(6))


def tim_bai_theo_ma(ma_bai):
    if db is None:
        return None

    r = (
        db.table("assignments")
        .select("*")
        .eq("code", ma_bai.strip().upper())
        .limit(1)
        .execute()
    )

    return r.data[0] if r.data else None


def tao_bai_moi(ten_bai, mo_ta, loai_file, barem):
    if db is None:
        raise ValueError("Chưa kết nối Supabase.")

    for _ in range(30):
        ma = tao_ma_ngan()
        if tim_bai_theo_ma(ma) is None:
            break
    else:
        raise ValueError("Không tạo được mã bài.")

    rubric = {
        "loai_file": loai_file,
        **barem
    }

    r = (
        db.table("assignments")
        .insert({
            "code": ma,
            "title": ten_bai,
            "description": mo_ta,
            "rubric": rubric,
            "active": True
        })
        .execute()
    )

    if not r.data:
        raise ValueError("Không lưu được bài tập.")

    return ma


def lay_danh_sach_bai():
    if db is None:
        return []

    r = (
        db.table("assignments")
        .select("id,code,title,description,active,created_at,rubric")
        .order("created_at", desc=True)
        .execute()
    )

    return r.data or []


def cap_nhat_trang_thai_bai(ma_bai, active):
    if db is None:
        return

    (
        db.table("assignments")
        .update({"active": bool(active)})
        .eq("code", ma_bai)
        .execute()
    )


def luu_ket_qua(ma_bai, ho_ten, lop, ten_file, ket_qua, nhan_xet):
    if db is None:
        return

    (
        db.table("submissions")
        .insert({
            "assignment_code": ma_bai,
            "student_name": ho_ten,
            "class_name": lop,
            "file_name": ten_file,
            "score": float(ket_qua["tong_diem"]),
            "details": ket_qua,
            "ai_feedback": nhan_xet
        })
        .execute()
    )


def lay_ket_qua_bai(ma_bai):
    if db is None:
        return []

    r = (
        db.table("submissions")
        .select(
            "id,assignment_code,student_name,class_name,"
            "file_name,score,submitted_at"
        )
        .eq("assignment_code", ma_bai)
        .order("submitted_at", desc=False)
        .execute()
    )

    return r.data or []


# =========================================================
# 4. HÀM CHUNG CHẤM ĐIỂM
# =========================================================

def ket_qua_tieu_chi(tc, dung, sai):
    tong = len(dung) + len(sai)

    if tong == 0:
        diem = 0
    else:
        diem = len(dung) / tong * float(tc["diem"])

    return {
        "mo_ta": tc["mo_ta"],
        "loai": tc["loai"],
        "diem": round(diem, 2),
        "diem_toi_da": float(tc["diem"]),
        "so_dung": len(dung),
        "tong": tong,
        "sai": sai
    }


def ket_qua_dat_khong(tc, dat, chi_tiet_dung="", chi_tiet_sai=""):
    return {
        "mo_ta": tc["mo_ta"],
        "loai": tc["loai"],
        "diem": float(tc["diem"]) if dat else 0,
        "diem_toi_da": float(tc["diem"]),
        "so_dung": 1 if dat else 0,
        "tong": 1,
        "sai": [] if dat else [chi_tiet_sai or "Chưa đạt"],
        "ghi_chu": chi_tiet_dung if dat else chi_tiet_sai
    }


# =========================================================
# 5. EXCEL
# =========================================================

def chuan_hoa_cong_thuc(x):
    if not isinstance(x, str):
        return ""

    x = x.replace("$", "").replace(";", ",")
    return re.sub(r"\s+", "", x).upper()


def co_du_4_vien(cell):
    b = cell.border
    if b is None:
        return False

    canh = [b.left, b.right, b.top, b.bottom]

    return all(
        x is not None and x.style is not None
        for x in canh
    )


def tim_hang_excel(sheet, barem):
    cot = barem["cot_moc"]
    start = int(barem["hang_bat_dau"])

    return [
        r
        for r in range(start, sheet.max_row + 1)
        if sheet[f"{cot}{r}"].value not in (None, "")
    ]


def cham_excel(file_bytes, barem):
    try:
        wb = openpyxl.load_workbook(
            io.BytesIO(file_bytes),
            data_only=False
        )
    except Exception as e:
        return {"loi": f"Không đọc được Excel: {e}"}

    ten_sheet = barem.get("ten_sheet", "").strip()

    if ten_sheet:
        if ten_sheet not in wb.sheetnames:
            return {"loi": f"Không tìm thấy sheet '{ten_sheet}'."}
        sh = wb[ten_sheet]
    else:
        sh = wb.active

    hang_dl = tim_hang_excel(sh, barem)

    if not hang_dl:
        return {"loi": "Không xác định được hàng dữ liệu."}

    kq = []

    for tc in barem["tieu_chi"]:
        loai = tc["loai"]
        dung, sai = [], []

        if loai == "HAM_VUNG":
            for r in hang_dl:
                o = f"{tc['cot_ket_qua']}{r}"

                thuc = chuan_hoa_cong_thuc(
                    sh[o].value
                )

                mong = chuan_hoa_cong_thuc(
                    f"={tc['ten_ham']}("
                    f"{tc['cot_dau']}{r}:"
                    f"{tc['cot_cuoi']}{r})"
                )

                if thuc == mong:
                    dung.append(o)
                else:
                    sai.append({
                        "o": o,
                        "hien_tai": sh[o].value
                    })

        elif loai == "CONG_THUC_MAU":
            for r in hang_dl:
                o = f"{tc['cot_ket_qua']}{r}"

                thuc = chuan_hoa_cong_thuc(
                    sh[o].value
                )

                mong = chuan_hoa_cong_thuc(
                    tc["cong_thuc_mau"].replace(
                        "{r}",
                        str(r)
                    )
                )

                if thuc == mong:
                    dung.append(o)
                else:
                    sai.append({
                        "o": o,
                        "hien_tai": sh[o].value
                    })

        elif loai == "IN_DAM_HANG":
            c1 = column_index_from_string(tc["cot_dau"])
            c2 = column_index_from_string(tc["cot_cuoi"])

            for c in range(c1, c2 + 1):
                cell = sh.cell(int(tc["hang"]), c)

                if cell.value in (None, ""):
                    continue

                if bool(cell.font.bold):
                    dung.append(cell.coordinate)
                else:
                    sai.append(cell.coordinate)

        elif loai == "VIEN_BANG":
            c1 = column_index_from_string(tc["cot_dau"])
            c2 = column_index_from_string(tc["cot_cuoi"])
            h2 = max(hang_dl)

            for r in range(int(tc["hang"]), h2 + 1):
                for c in range(c1, c2 + 1):
                    cell = sh.cell(r, c)

                    if co_du_4_vien(cell):
                        dung.append(cell.coordinate)
                    else:
                        sai.append(cell.coordinate)

            sai = sai[:50]

        elif loai == "CAN_LE_HANG":
            c1 = column_index_from_string(tc["cot_dau"])
            c2 = column_index_from_string(tc["cot_cuoi"])
            can_dung = tc["can_le"].lower()

            for c in range(c1, c2 + 1):
                cell = sh.cell(int(tc["hang"]), c)

                if cell.value in (None, ""):
                    continue

                can = (
                    cell.alignment.horizontal or ""
                ).lower()

                if can == can_dung:
                    dung.append(cell.coordinate)
                else:
                    sai.append(cell.coordinate)

        elif loai == "DINH_DANG_SO":
            for r in hang_dl:
                cell = sh[f"{tc['cot_ket_qua']}{r}"]

                if (
                    str(cell.number_format).strip()
                    == tc["dinh_dang_so"].strip()
                ):
                    dung.append(cell.coordinate)
                else:
                    sai.append(cell.coordinate)

        elif loai == "HAM_TRONG_COT":
            ten_ham = tc["ten_ham"].strip().upper()
            mau = re.compile(r"\\b" + re.escape(ten_ham) + r"\\s*\\(")

            for r in hang_dl:
                cell = sh[f"{tc['cot_ket_qua']}{r}"]
                cong_thuc = chuan_hoa_cong_thuc(cell.value)

                if (
                    cong_thuc.startswith("=")
                    and mau.search(cong_thuc)
                ):
                    dung.append(cell.coordinate)
                else:
                    sai.append({
                        "o": cell.coordinate,
                        "hien_tai": cell.value
                    })

        elif loai == "SO_BIEU_DO_TOI_THIEU":
            so = len(getattr(sh, "_charts", []))
            dat = so >= int(tc["gia_tri"])

            kq.append(
                ket_qua_dat_khong(
                    tc,
                    dat,
                    f"Có {so} biểu đồ.",
                    f"Hiện có {so} biểu đồ."
                )
            )
            continue

        elif loai == "CO_XAC_THUC_DU_LIEU":
            try:
                so = len(sh.data_validations.dataValidation)
            except Exception:
                so = 0

            dat = so >= int(tc.get("gia_tri", 1) or 1)

            kq.append(
                ket_qua_dat_khong(
                    tc,
                    dat,
                    f"Có {so} quy tắc xác thực dữ liệu.",
                    f"Hiện có {so} quy tắc xác thực dữ liệu."
                )
            )
            continue

        elif loai == "CO_BO_LOC":
            ref = getattr(sh.auto_filter, "ref", None)
            dat = bool(ref)

            kq.append(
                ket_qua_dat_khong(
                    tc,
                    dat,
                    f"Có vùng bộ lọc: {ref}.",
                    "Chưa phát hiện vùng bộ lọc."
                )
            )
            continue

        else:
            continue

        kq.append(
            ket_qua_tieu_chi(
                tc,
                dung,
                sai
            )
        )

    tong = round(
        sum(x["diem"] for x in kq),
        2
    )

    return {
        "loai_file": "Excel",
        "ten_bai": barem["ten_bai"],
        "tong_diem": min(tong, 10),
        "chi_tiet": kq
    }


# =========================================================
# 6. WORD
# =========================================================

def tat_ca_doan_word(doc):
    ket_qua = []

    for i, p in enumerate(doc.paragraphs, 1):
        if p.text.strip():
            ket_qua.append(
                (f"Đoạn {i}", p)
            )

    for ti, table in enumerate(doc.tables, 1):
        for ri, row in enumerate(table.rows, 1):
            for ci, cell in enumerate(row.cells, 1):
                for pi, p in enumerate(cell.paragraphs, 1):
                    if p.text.strip():
                        ket_qua.append(
                            (
                                f"Bảng {ti} "
                                f"ô[{ri},{ci}] "
                                f"đoạn {pi}",
                                p
                            )
                        )

    return ket_qua


def can_le_word(paragraph):
    a = paragraph.alignment

    if a is None:
        return "default"

    mapping = {
        WD_ALIGN_PARAGRAPH.LEFT: "left",
        WD_ALIGN_PARAGRAPH.CENTER: "center",
        WD_ALIGN_PARAGRAPH.RIGHT: "right",
        WD_ALIGN_PARAGRAPH.JUSTIFY: "justify"
    }

    return mapping.get(
        a,
        str(a).lower()
    )


def run_co_font(run, ten_font):
    if run.font.name:
        return (
            run.font.name.lower()
            == ten_font.lower()
        )

    return False


def run_co_size(run, co_chu):
    if run.font.size is None:
        return False

    return abs(
        run.font.size.pt
        - float(co_chu)
    ) < 0.2


def cham_word(file_bytes, barem):
    try:
        doc = Document(
            io.BytesIO(file_bytes)
        )
    except Exception as e:
        return {
            "loi":
                f"Không đọc được Word: {e}"
        }

    ds_doan = tat_ca_doan_word(doc)
    kq = []

    for tc in barem["tieu_chi"]:
        loai = tc["loai"]
        dung, sai = [], []

        if loai == "CO_NOI_DUNG":
            noi_dung = "\n".join(
                p.text
                for _, p in ds_doan
            ).lower()

            tu_khoa = [
                x.strip()
                for x in tc["gia_tri"].split("|")
                if x.strip()
            ]

            for x in tu_khoa:
                if x.lower() in noi_dung:
                    dung.append(x)
                else:
                    sai.append(x)

        elif loai == "SO_DOAN_TOI_THIEU":
            so = len([
                p
                for p in doc.paragraphs
                if p.text.strip()
            ])

            dat = so >= int(tc["gia_tri"])

            kq.append(
                ket_qua_dat_khong(
                    tc,
                    dat,
                    f"Có {so} đoạn.",
                    f"Hiện có {so} đoạn."
                )
            )
            continue

        elif loai == "SO_BANG_TOI_THIEU":
            so = len(doc.tables)
            dat = so >= int(tc["gia_tri"])

            kq.append(
                ket_qua_dat_khong(
                    tc,
                    dat,
                    f"Có {so} bảng.",
                    f"Hiện có {so} bảng."
                )
            )
            continue

        elif loai in (
            "IN_DAM",
            "IN_NGHIENG",
            "FONT",
            "CO_CHU",
            "CAN_LE"
        ):
            for ten, p in ds_doan:
                if loai == "CAN_LE":
                    ok = (
                        can_le_word(p)
                        == tc["gia_tri"].lower()
                    )
                else:
                    runs = [
                        r
                        for r in p.runs
                        if r.text.strip()
                    ]

                    if loai == "IN_DAM":
                        ok = bool(runs) and all(
                            bool(r.bold)
                            for r in runs
                        )

                    elif loai == "IN_NGHIENG":
                        ok = bool(runs) and all(
                            bool(r.italic)
                            for r in runs
                        )

                    elif loai == "FONT":
                        ok = bool(runs) and all(
                            run_co_font(
                                r,
                                tc["gia_tri"]
                            )
                            for r in runs
                        )

                    elif loai == "CO_CHU":
                        ok = bool(runs) and all(
                            run_co_size(
                                r,
                                tc["gia_tri"]
                            )
                            for r in runs
                        )

                if ok:
                    dung.append(ten)
                else:
                    sai.append(ten)

        elif loai == "GIAN_DONG":
            for ten, p in ds_doan:
                val = (
                    p.paragraph_format
                    .line_spacing
                )

                ok = False

                if val is not None:
                    try:
                        ok = abs(
                            float(val)
                            - float(tc["gia_tri"])
                        ) < 0.05
                    except Exception:
                        ok = False

                if ok:
                    dung.append(ten)
                else:
                    sai.append(ten)

        elif loai == "LE_TRANG":
            try:
                trai, phai, tren, duoi = [
                    float(x.strip())
                    for x
                    in tc["gia_tri"].split(",")
                ]
            except Exception:
                return {
                    "loi":
                        "Lề trang phải nhập "
                        "dạng 3,2,2,2"
                }

            for i, sec in enumerate(
                doc.sections,
                1
            ):
                vals = [
                    sec.left_margin.cm,
                    sec.right_margin.cm,
                    sec.top_margin.cm,
                    sec.bottom_margin.cm
                ]

                muc_tieu = [
                    trai,
                    phai,
                    tren,
                    duoi
                ]

                ok = all(
                    abs(a - b) < 0.15
                    for a, b
                    in zip(vals, muc_tieu)
                )

                if ok:
                    dung.append(
                        f"Section {i}"
                    )
                else:
                    sai.append(
                        f"Section {i}"
                    )

        elif loai == "SO_ANH_TOI_THIEU":
            so = len(doc.inline_shapes)
            dat = so >= int(tc["gia_tri"])

            kq.append(
                ket_qua_dat_khong(
                    tc,
                    dat,
                    f"Có {so} hình ảnh.",
                    f"Hiện có {so} hình ảnh."
                )
            )
            continue

        elif loai == "DAU_TRANG_CHAN_TRANG":
            co_noi_dung = False

            for sec in doc.sections:
                header_text = "\n".join(
                    p.text.strip()
                    for p in sec.header.paragraphs
                    if p.text.strip()
                )
                footer_text = "\n".join(
                    p.text.strip()
                    for p in sec.footer.paragraphs
                    if p.text.strip()
                )

                if header_text or footer_text:
                    co_noi_dung = True
                    break

            kq.append(
                ket_qua_dat_khong(
                    tc,
                    co_noi_dung,
                    "Có đầu trang hoặc chân trang.",
                    "Chưa phát hiện đầu trang/chân trang có nội dung."
                )
            )
            continue

        elif loai == "DANH_SACH_LIET_KE":
            so = 0

            for _, p in ds_doan:
                try:
                    pPr = p._p.pPr
                    if pPr is not None and pPr.numPr is not None:
                        so += 1
                        continue
                except Exception:
                    pass

                try:
                    style_name = p.style.name or ""
                    if "List" in style_name:
                        so += 1
                except Exception:
                    pass

            dat = so >= int(tc.get("gia_tri", 1) or 1)

            kq.append(
                ket_qua_dat_khong(
                    tc,
                    dat,
                    f"Có {so} đoạn danh sách.",
                    f"Hiện phát hiện {so} đoạn danh sách."
                )
            )
            continue

        else:
            continue

        kq.append(
            ket_qua_tieu_chi(
                tc,
                dung,
                sai[:50]
            )
        )

    tong = round(
        sum(x["diem"] for x in kq),
        2
    )

    return {
        "loai_file": "Word",
        "ten_bai": barem["ten_bai"],
        "tong_diem": min(tong, 10),
        "chi_tiet": kq
    }


# =========================================================
# 7. POWERPOINT - HÀM ĐỌC
# =========================================================

def ppt_tat_ca_shapes(shapes, prefix=""):
    """
    Trả về toàn bộ shape, kể cả shape trong group.
    """
    ket_qua = []

    for shape in shapes:
        ten = getattr(
            shape,
            "name",
            "Không tên"
        )

        path = (
            f"{prefix}/{ten}"
            if prefix
            else ten
        )

        if (
            shape.shape_type
            == MSO_SHAPE_TYPE.GROUP
        ):
            ket_qua.extend(
                ppt_tat_ca_shapes(
                    shape.shapes,
                    path
                )
            )
        else:
            ket_qua.append(
                (path, shape)
            )

    return ket_qua


def ppt_tat_ca_text_runs(prs):
    """
    Trả về (vị trí, paragraph, run) trên toàn bộ slide.
    """
    ket_qua = []

    for si, slide in enumerate(
        prs.slides,
        1
    ):
        for path, shape in ppt_tat_ca_shapes(
            slide.shapes
        ):
            if getattr(
                shape,
                "has_text_frame",
                False
            ):
                for pi, p in enumerate(
                    shape.text_frame.paragraphs,
                    1
                ):
                    for ri, run in enumerate(
                        p.runs,
                        1
                    ):
                        if run.text.strip():
                            ket_qua.append(
                                (
                                    f"Slide {si} "
                                    f"{path} "
                                    f"đoạn {pi} "
                                    f"run {ri}",
                                    p,
                                    run
                                )
                            )

            if getattr(
                shape,
                "has_table",
                False
            ):
                for rowi, row in enumerate(
                    shape.table.rows,
                    1
                ):
                    for coli, cell in enumerate(
                        row.cells,
                        1
                    ):
                        for pi, p in enumerate(
                            cell.text_frame.paragraphs,
                            1
                        ):
                            for ri, run in enumerate(
                                p.runs,
                                1
                            ):
                                if run.text.strip():
                                    ket_qua.append(
                                        (
                                            f"Slide {si} "
                                            f"bảng {path} "
                                            f"ô[{rowi},{coli}] "
                                            f"đoạn {pi} "
                                            f"run {ri}",
                                            p,
                                            run
                                        )
                                    )

    return ket_qua


def ppt_noi_dung_toan_bo(prs):
    parts = []

    for slide in prs.slides:
        for _, shape in ppt_tat_ca_shapes(
            slide.shapes
        ):
            if getattr(
                shape,
                "has_text_frame",
                False
            ):
                txt = shape.text.strip()
                if txt:
                    parts.append(txt)

            if getattr(
                shape,
                "has_table",
                False
            ):
                for row in shape.table.rows:
                    for cell in row.cells:
                        txt = cell.text.strip()
                        if txt:
                            parts.append(txt)

    return "\n".join(parts)


def ppt_dem_anh(prs):
    dem = 0

    for slide in prs.slides:
        for _, shape in ppt_tat_ca_shapes(
            slide.shapes
        ):
            if (
                shape.shape_type
                == MSO_SHAPE_TYPE.PICTURE
            ):
                dem += 1

    return dem


def ppt_dem_bang(prs):
    dem = 0

    for slide in prs.slides:
        for _, shape in ppt_tat_ca_shapes(
            slide.shapes
        ):
            if getattr(
                shape,
                "has_table",
                False
            ):
                dem += 1

    return dem


def ppt_so_slide_co_tieu_de(prs):
    dung = []
    sai = []

    for i, slide in enumerate(
        prs.slides,
        1
    ):
        title = slide.shapes.title

        if (
            title is not None
            and getattr(
                title,
                "has_text_frame",
                False
            )
            and title.text.strip()
        ):
            dung.append(
                f"Slide {i}"
            )
        else:
            sai.append(
                f"Slide {i}"
            )

    return dung, sai


def ppt_can_le(paragraph):
    a = paragraph.alignment

    if a is None:
        return "default"

    mapping = {
        PP_ALIGN.LEFT: "left",
        PP_ALIGN.CENTER: "center",
        PP_ALIGN.RIGHT: "right",
        PP_ALIGN.JUSTIFY: "justify"
    }

    return mapping.get(
        a,
        str(a).lower()
    )


# =========================================================
# 8. POWERPOINT - CHẤM
# =========================================================

def cham_powerpoint(file_bytes, barem):
    try:
        prs = Presentation(
            io.BytesIO(file_bytes)
        )
    except Exception as e:
        return {
            "loi":
                f"Không đọc được PowerPoint: {e}"
        }

    kq = []
    all_runs = ppt_tat_ca_text_runs(prs)

    for tc in barem["tieu_chi"]:
        loai = tc["loai"]

        if loai == "SO_SLIDE_TOI_THIEU":
            so = len(prs.slides)
            dat = so >= int(tc["gia_tri"])

            kq.append(
                ket_qua_dat_khong(
                    tc,
                    dat,
                    f"Có {so} slide.",
                    f"Hiện có {so} slide."
                )
            )
            continue

        elif loai == "CO_NOI_DUNG":
            noi_dung = (
                ppt_noi_dung_toan_bo(prs)
                .lower()
            )

            tu_khoa = [
                x.strip()
                for x in tc["gia_tri"].split("|")
                if x.strip()
            ]

            dung = []
            sai = []

            for x in tu_khoa:
                if x.lower() in noi_dung:
                    dung.append(x)
                else:
                    sai.append(x)

            kq.append(
                ket_qua_tieu_chi(
                    tc,
                    dung,
                    sai
                )
            )
            continue

        elif loai == "TIEU_DE_SLIDE":
            dung, sai = (
                ppt_so_slide_co_tieu_de(prs)
            )

            kq.append(
                ket_qua_tieu_chi(
                    tc,
                    dung,
                    sai
                )
            )
            continue

        elif loai == "SO_ANH_TOI_THIEU":
            so = ppt_dem_anh(prs)
            dat = so >= int(tc["gia_tri"])

            kq.append(
                ket_qua_dat_khong(
                    tc,
                    dat,
                    f"Có {so} hình ảnh.",
                    f"Hiện có {so} hình ảnh."
                )
            )
            continue

        elif loai == "SO_BANG_TOI_THIEU":
            so = ppt_dem_bang(prs)
            dat = so >= int(tc["gia_tri"])

            kq.append(
                ket_qua_dat_khong(
                    tc,
                    dat,
                    f"Có {so} bảng.",
                    f"Hiện có {so} bảng."
                )
            )
            continue

        elif loai in (
            "FONT",
            "CO_CHU",
            "IN_DAM",
            "IN_NGHIENG",
            "CAN_LE"
        ):
            dung = []
            sai = []

            for vi_tri, p, run in all_runs:
                if loai == "FONT":
                    ok = (
                        run.font.name is not None
                        and run.font.name.lower()
                        == tc["gia_tri"].lower()
                    )

                elif loai == "CO_CHU":
                    ok = (
                        run.font.size is not None
                        and abs(
                            run.font.size.pt
                            - float(tc["gia_tri"])
                        ) < 0.2
                    )

                elif loai == "IN_DAM":
                    ok = bool(
                        run.font.bold
                    )

                elif loai == "IN_NGHIENG":
                    ok = bool(
                        run.font.italic
                    )

                elif loai == "CAN_LE":
                    ok = (
                        ppt_can_le(p)
                        == tc["gia_tri"].lower()
                    )

                if ok:
                    dung.append(vi_tri)
                else:
                    sai.append(vi_tri)

            kq.append(
                ket_qua_tieu_chi(
                    tc,
                    dung,
                    sai[:50]
                )
            )
            continue

    tong = round(
        sum(x["diem"] for x in kq),
        2
    )

    return {
        "loai_file": "PowerPoint",
        "ten_bai": barem["ten_bai"],
        "tong_diem": min(tong, 10),
        "chi_tiet": kq
    }


# =========================================================
# 9. GEMINI NHẬN XÉT
# =========================================================

def tao_du_lieu_ai(ket_qua):
    dong = [
        f"Tên bài: {ket_qua['ten_bai']}",
        f"Loại bài: {ket_qua['loai_file']}",
        f"Điểm chính thức: {ket_qua['tong_diem']}/10"
    ]

    for tc in ket_qua["chi_tiet"]:
        dong.append(
            f"\n- {tc['mo_ta']}: "
            f"{tc['diem']}/"
            f"{tc['diem_toi_da']} điểm"
        )

        if tc["sai"]:
            loi = []

            for x in tc["sai"][:10]:
                if isinstance(x, dict):
                    loi.append(
                        str(x.get("o", x))
                    )
                else:
                    loi.append(str(x))

            dong.append(
                "Cần xem lại: "
                + ", ".join(loi)
            )

    return "\n".join(dong)


def tao_nhan_xet_ai(
    ho_ten,
    lop,
    ket_qua
):
    if not GEMINI_API_KEY:
        return ""

    try:
        client = genai.Client(
            api_key=GEMINI_API_KEY
        )

        prompt = f"""
Bạn là giáo viên Tin học THCS Việt Nam.

Học sinh: {ho_ten}
Lớp: {lop}

Kết quả chấm chính thức do Python xác định:
{tao_du_lieu_ai(ket_qua)}

QUY TẮC:
1. Tuyệt đối không thay đổi điểm.
2. Không tự chấm lại.
3. Không giải hộ bài.
4. Không viết sẵn công thức/đáp án hoàn chỉnh.
5. Chỉ rõ vị trí hoặc nội dung cần xem lại.
6. Đặt câu hỏi gợi mở kiểu Socratic.
7. Giọng văn khích lệ, chuẩn mực, ngắn gọn.

Trình bày:
🌟 **Điểm làm tốt**
🔍 **Nội dung cần xem lại**
💡 **Câu hỏi gợi mở**
🎯 **Điểm chính thức:** {ket_qua['tong_diem']}/10
"""

        r = client.models.generate_content(
            model=GEMINI_MODEL,
            contents=prompt,
            config=types.GenerateContentConfig(
                temperature=0.2
            )
        )

        return r.text or ""

    except Exception:
        return ""


# =========================================================
# 10. TẠO BAREM TỪ WEB
# =========================================================

def clean(v):
    if pd.isna(v):
        return None

    if (
        isinstance(v, float)
        and v.is_integer()
    ):
        return int(v)

    return v


def kiem_tra_tong_diem(tieu_chi):
    tong = sum(
        float(x["diem"])
        for x in tieu_chi
    )

    if abs(tong - 10) > 0.001:
        raise ValueError(
            f"Tổng điểm phải bằng 10. "
            f"Hiện tại: {tong:g}"
        )


def tao_barem_excel(
    ten_bai,
    ten_sheet,
    cot_moc,
    hang_bat_dau,
    df
):
    tieu_chi = []

    for row in df.to_dict(
        orient="records"
    ):
        row = {
            k: clean(v)
            for k, v in row.items()
        }

        loai = row.get("loai")
        mo_ta = row.get("mo_ta")
        diem = float(
            row.get("diem") or 0
        )

        if not loai or not mo_ta or diem <= 0:
            continue

        tc = {
            "loai": str(loai),
            "mo_ta": str(mo_ta),
            "diem": diem
        }

        if loai == "HAM_VUNG":
            tc.update({
                "cot_ket_qua":
                    str(
                        row.get(
                            "cot_ket_qua"
                        ) or ""
                    ).upper(),

                "ten_ham":
                    str(
                        row.get(
                            "ten_ham"
                        ) or ""
                    ).upper(),

                "cot_dau":
                    str(
                        row.get(
                            "cot_dau"
                        ) or ""
                    ).upper(),

                "cot_cuoi":
                    str(
                        row.get(
                            "cot_cuoi"
                        ) or ""
                    ).upper()
            })

        elif loai == "CONG_THUC_MAU":
            tc.update({
                "cot_ket_qua":
                    str(
                        row.get(
                            "cot_ket_qua"
                        ) or ""
                    ).upper(),

                "cong_thuc_mau":
                    str(
                        row.get(
                            "cong_thuc_mau"
                        ) or ""
                    )
            })

        elif loai in (
            "IN_DAM_HANG",
            "VIEN_BANG",
            "CAN_LE_HANG"
        ):
            tc.update({
                "hang":
                    int(
                        row.get("hang")
                        or 1
                    ),

                "cot_dau":
                    str(
                        row.get(
                            "cot_dau"
                        ) or "A"
                    ).upper(),

                "cot_cuoi":
                    str(
                        row.get(
                            "cot_cuoi"
                        ) or "A"
                    ).upper()
            })

            if loai == "CAN_LE_HANG":
                tc["can_le"] = str(
                    row.get("can_le")
                    or "center"
                ).lower()

        elif loai == "DINH_DANG_SO":
            tc.update({
                "cot_ket_qua":
                    str(
                        row.get(
                            "cot_ket_qua"
                        ) or ""
                    ).upper(),

                "dinh_dang_so":
                    str(
                        row.get(
                            "dinh_dang_so"
                        ) or ""
                    )
            })

        elif loai == "HAM_TRONG_COT":
            tc.update({
                "cot_ket_qua":
                    str(
                        row.get(
                            "cot_ket_qua"
                        ) or ""
                    ).upper(),

                "ten_ham":
                    str(
                        row.get(
                            "ten_ham"
                        ) or ""
                    ).upper()
            })

        elif loai in (
            "SO_BIEU_DO_TOI_THIEU",
            "CO_XAC_THUC_DU_LIEU",
            "CO_BO_LOC"
        ):
            tc["gia_tri"] = str(
                row.get("gia_tri") or "1"
            )

        tieu_chi.append(tc)

    kiem_tra_tong_diem(
        tieu_chi
    )

    return {
        "ten_bai": ten_bai,
        "ten_sheet": ten_sheet,
        "cot_moc": cot_moc.upper(),
        "hang_bat_dau": int(
            hang_bat_dau
        ),
        "tieu_chi": tieu_chi
    }


def tao_barem_dang_gia_tri(
    ten_bai,
    df
):
    tieu_chi = []

    for row in df.to_dict(
        orient="records"
    ):
        row = {
            k: clean(v)
            for k, v in row.items()
        }

        loai = row.get("loai")
        mo_ta = row.get("mo_ta")
        diem = float(
            row.get("diem") or 0
        )

        if not loai or not mo_ta or diem <= 0:
            continue

        tieu_chi.append({
            "loai": str(loai),
            "mo_ta": str(mo_ta),
            "diem": diem,
            "gia_tri":
                str(
                    row.get(
                        "gia_tri"
                    ) or ""
                ).strip()
        })

    kiem_tra_tong_diem(
        tieu_chi
    )

    return {
        "ten_bai": ten_bai,
        "tieu_chi": tieu_chi
    }


# =========================================================
# 11. XUẤT EXCEL KẾT QUẢ
# =========================================================

def tao_file_excel_ket_qua(ds):
    if not ds:
        return None

    df = pd.DataFrame(ds)

    df = df.rename(
        columns={
            "student_name":
                "Họ và tên",
            "class_name":
                "Lớp",
            "score":
                "Điểm",
            "file_name":
                "Tên file",
            "submitted_at":
                "Thời gian nộp"
        }
    )

    cols = [
        c
        for c in [
            "Họ và tên",
            "Lớp",
            "Điểm",
            "Tên file",
            "Thời gian nộp"
        ]
        if c in df.columns
    ]

    df = df[cols]

    out = io.BytesIO()

    with pd.ExcelWriter(
        out,
        engine="openpyxl"
    ) as writer:
        df.to_excel(
            writer,
            index=False,
            sheet_name="Ket_qua"
        )

        ws = writer.book[
            "Ket_qua"
        ]

        for cell in ws[1]:
            cell.font = (
                openpyxl.styles
                .Font(bold=True)
            )

        for col in ws.columns:
            width = min(
                max(
                    len(
                        str(
                            c.value or ""
                        )
                    )
                    for c in col
                ) + 2,
                35
            )

            ws.column_dimensions[
                col[0].column_letter
            ].width = width

    out.seek(0)
    return out.getvalue()



# =========================================================
# 11B. BAREM GỢI Ý THEO CTGDPT 2018 / SGK KẾT NỐI TRI THỨC
# =========================================================

def df_excel(rows):
    cot = [
        "loai", "mo_ta", "diem", "cot_ket_qua",
        "ten_ham", "cot_dau", "cot_cuoi", "hang",
        "cong_thuc_mau", "can_le", "dinh_dang_so",
        "gia_tri"
    ]
    return pd.DataFrame(rows, columns=cot)


def df_gia_tri(rows):
    return pd.DataFrame(
        rows,
        columns=["loai", "mo_ta", "diem", "gia_tri"]
    )


BAREM_CHUAN_EXCEL = {
    "Lớp 7 – Bảng tính cơ bản, công thức và hàm": [
        {
            "loai": "HAM_TRONG_COT",
            "mo_ta": "Sử dụng công thức/hàm tính toán đúng yêu cầu",
            "diem": 4.0,
            "cot_ket_qua": "F",
            "ten_ham": "AVERAGE"
        },
        {
            "loai": "DINH_DANG_SO",
            "mo_ta": "Định dạng dữ liệu số phù hợp",
            "diem": 2.0,
            "cot_ket_qua": "F",
            "dinh_dang_so": "0.00"
        },
        {
            "loai": "IN_DAM_HANG",
            "mo_ta": "Hàng tiêu đề được định dạng nổi bật",
            "diem": 2.0,
            "hang": 1,
            "cot_dau": "A",
            "cot_cuoi": "F"
        },
        {
            "loai": "VIEN_BANG",
            "mo_ta": "Bảng dữ liệu được kẻ viền đầy đủ",
            "diem": 2.0,
            "hang": 1,
            "cot_dau": "A",
            "cot_cuoi": "F"
        }
    ],

    "Lớp 7 – Hoàn thiện bảng tính": [
        {
            "loai": "HAM_TRONG_COT",
            "mo_ta": "Công thức/hàm tính toán đúng",
            "diem": 4.0,
            "cot_ket_qua": "F",
            "ten_ham": "AVERAGE"
        },
        {
            "loai": "DINH_DANG_SO",
            "mo_ta": "Định dạng số đúng yêu cầu",
            "diem": 2.0,
            "cot_ket_qua": "F",
            "dinh_dang_so": "0.00"
        },
        {
            "loai": "VIEN_BANG",
            "mo_ta": "Kẻ khung vùng dữ liệu",
            "diem": 2.0,
            "hang": 1,
            "cot_dau": "A",
            "cot_cuoi": "F"
        },
        {
            "loai": "IN_DAM_HANG",
            "mo_ta": "Tiêu đề bảng rõ ràng, nổi bật",
            "diem": 2.0,
            "hang": 1,
            "cot_dau": "A",
            "cot_cuoi": "F"
        }
    ],

    "Lớp 8 – Sử dụng bảng tính giải quyết bài toán thực tế": [
        {
            "loai": "HAM_TRONG_COT",
            "mo_ta": "Sử dụng công thức/hàm phù hợp để giải quyết bài toán",
            "diem": 5.0,
            "cot_ket_qua": "F",
            "ten_ham": "SUM"
        },
        {
            "loai": "DINH_DANG_SO",
            "mo_ta": "Định dạng số/dữ liệu phù hợp",
            "diem": 2.0,
            "cot_ket_qua": "F",
            "dinh_dang_so": "#,##0"
        },
        {
            "loai": "VIEN_BANG",
            "mo_ta": "Trình bày bảng rõ ràng, có đường viền",
            "diem": 1.5,
            "hang": 1,
            "cot_dau": "A",
            "cot_cuoi": "F"
        },
        {
            "loai": "IN_DAM_HANG",
            "mo_ta": "Tiêu đề bảng được định dạng phù hợp",
            "diem": 1.5,
            "hang": 1,
            "cot_dau": "A",
            "cot_cuoi": "F"
        }
    ],

    "Lớp 8 – Sắp xếp và lọc dữ liệu": [
        {
            "loai": "CO_BO_LOC",
            "mo_ta": "Thiết lập bộ lọc cho vùng dữ liệu",
            "diem": 5.0,
            "gia_tri": "1"
        },
        {
            "loai": "VIEN_BANG",
            "mo_ta": "Bảng dữ liệu được trình bày rõ ràng",
            "diem": 2.5,
            "hang": 1,
            "cot_dau": "A",
            "cot_cuoi": "F"
        },
        {
            "loai": "IN_DAM_HANG",
            "mo_ta": "Hàng tiêu đề rõ ràng",
            "diem": 2.5,
            "hang": 1,
            "cot_dau": "A",
            "cot_cuoi": "F"
        }
    ],

    "Lớp 8 – Trình bày dữ liệu bằng biểu đồ": [
        {
            "loai": "SO_BIEU_DO_TOI_THIEU",
            "mo_ta": "Tạo được biểu đồ biểu diễn dữ liệu",
            "diem": 6.0,
            "gia_tri": "1"
        },
        {
            "loai": "VIEN_BANG",
            "mo_ta": "Nguồn dữ liệu được trình bày rõ ràng",
            "diem": 2.0,
            "hang": 1,
            "cot_dau": "A",
            "cot_cuoi": "F"
        },
        {
            "loai": "IN_DAM_HANG",
            "mo_ta": "Tiêu đề nguồn dữ liệu được làm nổi bật",
            "diem": 2.0,
            "hang": 1,
            "cot_dau": "A",
            "cot_cuoi": "F"
        }
    ],

    "Lớp 9 – Sử dụng công cụ xác thực dữ liệu": [
        {
            "loai": "CO_XAC_THUC_DU_LIEU",
            "mo_ta": "Thiết lập xác thực dữ liệu theo yêu cầu",
            "diem": 7.0,
            "gia_tri": "1"
        },
        {
            "loai": "VIEN_BANG",
            "mo_ta": "Bảng dữ liệu được trình bày rõ ràng",
            "diem": 1.5,
            "hang": 1,
            "cot_dau": "A",
            "cot_cuoi": "F"
        },
        {
            "loai": "IN_DAM_HANG",
            "mo_ta": "Hàng tiêu đề được định dạng phù hợp",
            "diem": 1.5,
            "hang": 1,
            "cot_dau": "A",
            "cot_cuoi": "F"
        }
    ],

    "Lớp 9 – Hàm COUNTIF": [
        {
            "loai": "HAM_TRONG_COT",
            "mo_ta": "Sử dụng đúng hàm COUNTIF",
            "diem": 7.0,
            "cot_ket_qua": "F",
            "ten_ham": "COUNTIF"
        },
        {
            "loai": "VIEN_BANG",
            "mo_ta": "Bảng trình bày rõ ràng",
            "diem": 1.5,
            "hang": 1,
            "cot_dau": "A",
            "cot_cuoi": "F"
        },
        {
            "loai": "IN_DAM_HANG",
            "mo_ta": "Tiêu đề bảng phù hợp",
            "diem": 1.5,
            "hang": 1,
            "cot_dau": "A",
            "cot_cuoi": "F"
        }
    ],

    "Lớp 9 – Hàm SUMIF": [
        {
            "loai": "HAM_TRONG_COT",
            "mo_ta": "Sử dụng đúng hàm SUMIF",
            "diem": 7.0,
            "cot_ket_qua": "F",
            "ten_ham": "SUMIF"
        },
        {
            "loai": "DINH_DANG_SO",
            "mo_ta": "Định dạng số phù hợp",
            "diem": 1.5,
            "cot_ket_qua": "F",
            "dinh_dang_so": "#,##0"
        },
        {
            "loai": "VIEN_BANG",
            "mo_ta": "Bảng trình bày rõ ràng",
            "diem": 1.5,
            "hang": 1,
            "cot_dau": "A",
            "cot_cuoi": "F"
        }
    ],

    "Lớp 9 – Hàm IF": [
        {
            "loai": "HAM_TRONG_COT",
            "mo_ta": "Sử dụng đúng hàm IF",
            "diem": 7.0,
            "cot_ket_qua": "G",
            "ten_ham": "IF"
        },
        {
            "loai": "VIEN_BANG",
            "mo_ta": "Bảng trình bày rõ ràng",
            "diem": 1.5,
            "hang": 1,
            "cot_dau": "A",
            "cot_cuoi": "G"
        },
        {
            "loai": "IN_DAM_HANG",
            "mo_ta": "Tiêu đề bảng phù hợp",
            "diem": 1.5,
            "hang": 1,
            "cot_dau": "A",
            "cot_cuoi": "G"
        }
    ],

    "Lớp 9 – Hoàn thiện bảng tính quản lí tài chính": [
        {
            "loai": "HAM_TRONG_COT",
            "mo_ta": "Sử dụng hàm điều kiện/thống kê đúng yêu cầu",
            "diem": 5.0,
            "cot_ket_qua": "F",
            "ten_ham": "SUMIF"
        },
        {
            "loai": "CO_XAC_THUC_DU_LIEU",
            "mo_ta": "Có xác thực dữ liệu ở vùng cần nhập",
            "diem": 2.0,
            "gia_tri": "1"
        },
        {
            "loai": "DINH_DANG_SO",
            "mo_ta": "Định dạng tiền tệ/số phù hợp",
            "diem": 1.5,
            "cot_ket_qua": "F",
            "dinh_dang_so": "#,##0"
        },
        {
            "loai": "VIEN_BANG",
            "mo_ta": "Bảng quản lí được trình bày rõ ràng",
            "diem": 1.5,
            "hang": 1,
            "cot_dau": "A",
            "cot_cuoi": "G"
        }
    ]
}


BAREM_CHUAN_WORD = {
    "Lớp 6 – Định dạng văn bản": [
        {
            "loai": "FONT",
            "mo_ta": "Sử dụng font chữ theo yêu cầu",
            "diem": 2.0,
            "gia_tri": "Times New Roman"
        },
        {
            "loai": "CO_CHU",
            "mo_ta": "Cỡ chữ phù hợp yêu cầu",
            "diem": 2.0,
            "gia_tri": "14"
        },
        {
            "loai": "CAN_LE",
            "mo_ta": "Căn lề đoạn văn phù hợp",
            "diem": 2.0,
            "gia_tri": "justify"
        },
        {
            "loai": "IN_DAM",
            "mo_ta": "Thực hiện định dạng chữ đậm ở nội dung yêu cầu",
            "diem": 2.0,
            "gia_tri": ""
        },
        {
            "loai": "CO_NOI_DUNG",
            "mo_ta": "Có đầy đủ nội dung bắt buộc",
            "diem": 2.0,
            "gia_tri": "Họ và tên|Lớp"
        }
    ],

    "Lớp 6 – Trình bày thông tin ở dạng bảng": [
        {
            "loai": "SO_BANG_TOI_THIEU",
            "mo_ta": "Tạo được bảng để trình bày thông tin",
            "diem": 5.0,
            "gia_tri": "1"
        },
        {
            "loai": "CO_NOI_DUNG",
            "mo_ta": "Bảng có nội dung theo yêu cầu",
            "diem": 2.0,
            "gia_tri": "Họ và tên|Lớp"
        },
        {
            "loai": "FONT",
            "mo_ta": "Font chữ thống nhất, dễ đọc",
            "diem": 1.5,
            "gia_tri": "Times New Roman"
        },
        {
            "loai": "CO_CHU",
            "mo_ta": "Cỡ chữ phù hợp",
            "diem": 1.5,
            "gia_tri": "14"
        }
    ],

    "Lớp 6 – Thực hành tổng hợp: Sổ lưu niệm": [
        {
            "loai": "CO_NOI_DUNG",
            "mo_ta": "Có các nội dung chính của sản phẩm",
            "diem": 3.0,
            "gia_tri": "Họ và tên|Lớp"
        },
        {
            "loai": "SO_BANG_TOI_THIEU",
            "mo_ta": "Có bảng trình bày thông tin",
            "diem": 2.0,
            "gia_tri": "1"
        },
        {
            "loai": "SO_ANH_TOI_THIEU",
            "mo_ta": "Có hình ảnh minh hoạ",
            "diem": 2.0,
            "gia_tri": "1"
        },
        {
            "loai": "FONT",
            "mo_ta": "Định dạng font chữ thống nhất",
            "diem": 1.5,
            "gia_tri": "Times New Roman"
        },
        {
            "loai": "CAN_LE",
            "mo_ta": "Căn lề đoạn văn phù hợp",
            "diem": 1.5,
            "gia_tri": "justify"
        }
    ],

    "Lớp 8 – Danh sách liệt kê và hình ảnh trong văn bản": [
        {
            "loai": "DANH_SACH_LIET_KE",
            "mo_ta": "Tạo được danh sách dạng liệt kê",
            "diem": 4.0,
            "gia_tri": "1"
        },
        {
            "loai": "SO_ANH_TOI_THIEU",
            "mo_ta": "Chèn được hình ảnh minh hoạ",
            "diem": 3.0,
            "gia_tri": "1"
        },
        {
            "loai": "CO_NOI_DUNG",
            "mo_ta": "Nội dung văn bản đầy đủ",
            "diem": 2.0,
            "gia_tri": "Họ và tên|Lớp"
        },
        {
            "loai": "CAN_LE",
            "mo_ta": "Trình bày đoạn văn hợp lí",
            "diem": 1.0,
            "gia_tri": "justify"
        }
    ],

    "Lớp 8 – Tạo đầu trang, chân trang": [
        {
            "loai": "DAU_TRANG_CHAN_TRANG",
            "mo_ta": "Tạo được đầu trang hoặc chân trang",
            "diem": 6.0,
            "gia_tri": ""
        },
        {
            "loai": "FONT",
            "mo_ta": "Font chữ phù hợp, thống nhất",
            "diem": 2.0,
            "gia_tri": "Times New Roman"
        },
        {
            "loai": "CAN_LE",
            "mo_ta": "Nội dung văn bản trình bày hợp lí",
            "diem": 2.0,
            "gia_tri": "justify"
        }
    ]
}


BAREM_CHUAN_PPT = {
    "Lớp 8 – Định dạng nâng cao cho trang chiếu": [
        {
            "loai": "SO_SLIDE_TOI_THIEU",
            "mo_ta": "Có đủ số trang chiếu theo yêu cầu",
            "diem": 2.0,
            "gia_tri": "4"
        },
        {
            "loai": "TIEU_DE_SLIDE",
            "mo_ta": "Các trang chiếu có tiêu đề rõ ràng",
            "diem": 2.0,
            "gia_tri": ""
        },
        {
            "loai": "CO_NOI_DUNG",
            "mo_ta": "Nội dung chính đầy đủ, ngắn gọn",
            "diem": 2.0,
            "gia_tri": "Mở đầu|Nội dung|Kết luận"
        },
        {
            "loai": "SO_ANH_TOI_THIEU",
            "mo_ta": "Có hình ảnh minh hoạ",
            "diem": 2.0,
            "gia_tri": "1"
        },
        {
            "loai": "CO_CHU",
            "mo_ta": "Cỡ chữ đủ lớn để trình chiếu",
            "diem": 2.0,
            "gia_tri": "24"
        }
    ],

    "Lớp 8 – Sử dụng bản mẫu tạo bài trình chiếu": [
        {
            "loai": "SO_SLIDE_TOI_THIEU",
            "mo_ta": "Tạo được bài trình chiếu đủ số slide",
            "diem": 2.0,
            "gia_tri": "5"
        },
        {
            "loai": "TIEU_DE_SLIDE",
            "mo_ta": "Các slide có tiêu đề",
            "diem": 2.0,
            "gia_tri": ""
        },
        {
            "loai": "CO_NOI_DUNG",
            "mo_ta": "Nội dung phục vụ học tập/giao lưu đầy đủ",
            "diem": 2.0,
            "gia_tri": "Mở đầu|Nội dung|Kết luận"
        },
        {
            "loai": "SO_ANH_TOI_THIEU",
            "mo_ta": "Có hình ảnh minh hoạ phù hợp",
            "diem": 2.0,
            "gia_tri": "2"
        },
        {
            "loai": "CO_CHU",
            "mo_ta": "Văn bản trên slide có cỡ chữ dễ đọc",
            "diem": 2.0,
            "gia_tri": "24"
        }
    ],

    "Mẫu chung – Bài trình chiếu THCS": [
        {
            "loai": "SO_SLIDE_TOI_THIEU",
            "mo_ta": "Đủ số slide tối thiểu",
            "diem": 2.0,
            "gia_tri": "5"
        },
        {
            "loai": "TIEU_DE_SLIDE",
            "mo_ta": "Mỗi slide có tiêu đề khi phù hợp",
            "diem": 2.0,
            "gia_tri": ""
        },
        {
            "loai": "CO_NOI_DUNG",
            "mo_ta": "Có đủ nội dung bắt buộc",
            "diem": 2.0,
            "gia_tri": "Mở đầu|Nội dung|Kết luận"
        },
        {
            "loai": "SO_ANH_TOI_THIEU",
            "mo_ta": "Có hình ảnh minh hoạ",
            "diem": 2.0,
            "gia_tri": "2"
        },
        {
            "loai": "CO_CHU",
            "mo_ta": "Cỡ chữ dễ quan sát",
            "diem": 2.0,
            "gia_tri": "24"
        }
    ]
}


def lay_barem_goi_y(loai_file, ten_mau):
    if loai_file == "Excel":
        return df_excel(
            BAREM_CHUAN_EXCEL[ten_mau]
        )

    if loai_file == "Word":
        return df_gia_tri(
            BAREM_CHUAN_WORD[ten_mau]
        )

    return df_gia_tri(
        BAREM_CHUAN_PPT[ten_mau]
    )


# =========================================================
# 12. KIỂM TRA CẤU HÌNH
# =========================================================

thieu = []

if not TEACHER_PIN:
    thieu.append("TEACHER_PIN")

if not SUPABASE_URL:
    thieu.append("SUPABASE_URL")

if not SUPABASE_KEY:
    thieu.append(
        "SUPABASE_SERVICE_KEY"
    )

if thieu:
    st.warning(
        "Thiếu cấu hình: "
        + ", ".join(thieu)
    )


# =========================================================
# 13. GIAO DIỆN CHÍNH
# =========================================================

tab_hs, tab_gv = st.tabs([
    "🎓 HỌC SINH",
    "👩‍🏫 GIÁO VIÊN"
])


# =========================================================
# 14. HỌC SINH
# =========================================================

with tab_hs:
    st.subheader(
        "🎓 Nộp bài"
    )

    ma_bai = st.text_input(
        "Mã bài",
        placeholder="Ví dụ: AI-7K4P9Q"
    )

    ho_ten = st.text_input(
        "Họ và tên"
    )

    lop = st.text_input(
        "Lớp",
        placeholder="Ví dụ: 7A1"
    )

    bai = None
    file_nop = None

    if (
        ma_bai.strip()
        and db is not None
    ):
        bai = tim_bai_theo_ma(
            ma_bai.strip().upper()
        )

        if bai:
            loai_file = (
                bai["rubric"]
                .get(
                    "loai_file",
                    "Excel"
                )
            )

            st.info(
                f"📘 {bai['title']} — "
                f"Loại bài: {loai_file}"
            )

            if loai_file == "Word":
                file_nop = (
                    st.file_uploader(
                        "Chọn file Word",
                        type=["docx"]
                    )
                )

            elif loai_file == "PowerPoint":
                file_nop = (
                    st.file_uploader(
                        "Chọn file PowerPoint",
                        type=["pptx"]
                    )
                )

            else:
                file_nop = (
                    st.file_uploader(
                        "Chọn file Excel",
                        type=["xlsx"]
                    )
                )

    if st.button(
        "🚀 Nộp bài và chấm điểm",
        type="primary",
        use_container_width=True
    ):
        if db is None:
            st.error(
                "Chưa kết nối cơ sở dữ liệu."
            )

        elif not ma_bai.strip():
            st.warning(
                "Em hãy nhập mã bài."
            )

        elif not bai:
            st.error(
                "Không tìm thấy mã bài."
            )

        elif not bai["active"]:
            st.warning(
                "Bài này đã đóng nhận bài."
            )

        elif not ho_ten.strip():
            st.warning(
                "Em hãy nhập họ và tên."
            )

        elif not lop.strip():
            st.warning(
                "Em hãy nhập lớp."
            )

        elif file_nop is None:
            st.warning(
                "Em hãy chọn file bài làm."
            )

        else:
            barem = bai["rubric"]
            loai_file = barem.get(
                "loai_file",
                "Excel"
            )

            with st.spinner(
                "Đang chấm bài..."
            ):
                if loai_file == "Word":
                    ket_qua = cham_word(
                        file_nop.getvalue(),
                        barem
                    )

                elif loai_file == "PowerPoint":
                    ket_qua = (
                        cham_powerpoint(
                            file_nop.getvalue(),
                            barem
                        )
                    )

                else:
                    ket_qua = cham_excel(
                        file_nop.getvalue(),
                        barem
                    )

            if "loi" in ket_qua:
                st.error(
                    ket_qua["loi"]
                )

            else:
                with st.spinner(
                    "Đang tạo nhận xét..."
                ):
                    nhan_xet = (
                        tao_nhan_xet_ai(
                            ho_ten,
                            lop,
                            ket_qua
                        )
                    )

                try:
                    luu_ket_qua(
                        ma_bai
                        .strip()
                        .upper(),

                        ho_ten.strip(),
                        lop.strip(),
                        file_nop.name,
                        ket_qua,
                        nhan_xet
                    )

                except Exception as e:
                    st.warning(
                        "Đã chấm nhưng "
                        "chưa lưu được kết quả."
                    )
                    st.caption(
                        str(e)
                    )

                st.success(
                    "✅ Đã chấm xong!"
                )

                st.metric(
                    "🎯 Điểm",
                    f"{ket_qua['tong_diem']}/10"
                )

                st.markdown(
                    "### 📊 Chi tiết"
                )

                for tc in ket_qua[
                    "chi_tiet"
                ]:
                    icon = (
                        "✅"
                        if not tc["sai"]
                        else "⚠️"
                    )

                    st.markdown(
                        f"{icon} "
                        f"**{tc['mo_ta']}**"
                    )

                    st.write(
                        f"{tc['diem']}/"
                        f"{tc['diem_toi_da']} điểm"
                    )

                    if tc["sai"]:
                        vals = []

                        for x in tc["sai"][:20]:
                            if isinstance(
                                x,
                                dict
                            ):
                                vals.append(
                                    str(
                                        x.get(
                                            "o",
                                            x
                                        )
                                    )
                                )
                            else:
                                vals.append(
                                    str(x)
                                )

                        st.caption(
                            "Cần xem lại: "
                            + ", ".join(vals)
                        )

                if nhan_xet:
                    st.markdown(
                        "### 🤖 Nhận xét sư phạm"
                    )

                    st.markdown(
                        nhan_xet
                    )


# =========================================================
# 15. GIÁO VIÊN - ĐĂNG NHẬP
# =========================================================

with tab_gv:
    st.subheader(
        "👩‍🏫 Khu vực giáo viên"
    )

    if "gv_login" not in st.session_state:
        st.session_state[
            "gv_login"
        ] = False

    if not st.session_state[
        "gv_login"
    ]:
        pin = st.text_input(
            "Mật khẩu giáo viên",
            type="password"
        )

        if st.button(
            "🔐 Đăng nhập"
        ):
            if (
                TEACHER_PIN
                and hmac.compare_digest(
                    pin,
                    TEACHER_PIN
                )
            ):
                st.session_state[
                    "gv_login"
                ] = True

                st.rerun()

            else:
                st.error(
                    "Mật khẩu không đúng."
                )

    else:
        if st.button(
            "Đăng xuất"
        ):
            st.session_state[
                "gv_login"
            ] = False

            st.rerun()

        tao_tab, ql_tab = st.tabs([
            "➕ Tạo bài",
            "📊 Quản lý & kết quả"
        ])


        # =================================================
        # 16. GIÁO VIÊN - TẠO BÀI
        # =================================================

        with tao_tab:
            ten_bai = st.text_input(
                "Tên bài",
                value="Bài thực hành"
            )

            mo_ta = st.text_area(
                "Mô tả / yêu cầu chung"
            )

            loai_file = st.selectbox(
                "Loại bài học sinh nộp",
                [
                    "Excel",
                    "Word",
                    "PowerPoint"
                ]
            )

            st.info(
                "Barem chuẩn gợi ý được xây dựng theo yêu cầu cần đạt CTGDPT 2018 "
                "và mạch bài thực hành SGK Tin học THCS (tham chiếu Kết nối tri thức). "
                "Giáo viên được phép chỉnh lại theo đúng đề giao cho học sinh."
            )

            # ---------------------------------------------
            # EXCEL
            # ---------------------------------------------
            if loai_file == "Excel":
                c1, c2 = st.columns(2)

                with c1:
                    ten_sheet = st.text_input(
                        "Tên sheet cần chấm",
                        placeholder=(
                            "Để trống = "
                            "sheet đầu tiên"
                        )
                    )

                    cot_moc = st.text_input(
                        "Cột xác định hàng dữ liệu",
                        value="A"
                    )

                with c2:
                    hang_bat_dau = (
                        st.number_input(
                            "Dữ liệu bắt đầu từ hàng",
                            min_value=1,
                            value=2,
                            step=1
                        )
                    )

                st.markdown(
                    "### Barem Excel"
                )

                mau_excel = st.selectbox(
                    "Barem chuẩn gợi ý",
                    list(BAREM_CHUAN_EXCEL.keys()),
                    key="mau_excel"
                )

                st.caption(
                    "Barem được nạp theo yêu cầu cần đạt gợi ý. "
                    "Cô có thể sửa cột, vùng dữ liệu, hàm và điểm trước khi tạo bài."
                )

                mac_dinh = lay_barem_goi_y(
                    "Excel",
                    mau_excel
                )

                df = st.data_editor(
                    mac_dinh,
                    num_rows="dynamic",
                    use_container_width=True,
                    column_config={
                        "loai":
                            st.column_config
                            .SelectboxColumn(
                                "Loại tiêu chí",
                                options=[
                                    "HAM_VUNG",
                                    "CONG_THUC_MAU",
                                    "IN_DAM_HANG",
                                    "VIEN_BANG",
                                    "CAN_LE_HANG",
                                    "DINH_DANG_SO",
                                    "HAM_TRONG_COT",
                                    "SO_BIEU_DO_TOI_THIEU",
                                    "CO_XAC_THUC_DU_LIEU",
                                    "CO_BO_LOC"
                                ]
                            ),

                        "diem":
                            st.column_config
                            .NumberColumn(
                                "Điểm",
                                min_value=0.0,
                                max_value=10.0,
                                step=0.25
                            )
                    }
                )

                if st.button(
                    "✅ Tạo bài Excel và sinh mã",
                    type="primary"
                ):
                    try:
                        barem = (
                            tao_barem_excel(
                                ten_bai,
                                ten_sheet,
                                cot_moc,
                                hang_bat_dau,
                                df
                            )
                        )

                        ma = tao_bai_moi(
                            ten_bai,
                            mo_ta,
                            "Excel",
                            barem
                        )

                        st.success(
                            "Đã tạo bài!"
                        )
                        st.code(ma)

                    except Exception as e:
                        st.error(str(e))

            # ---------------------------------------------
            # WORD
            # ---------------------------------------------
            elif loai_file == "Word":
                st.markdown(
                    "### Barem Word"
                )

                mau_word = st.selectbox(
                    "Barem chuẩn gợi ý",
                    list(BAREM_CHUAN_WORD.keys()),
                    key="mau_word"
                )

                st.caption(
                    "Barem gợi ý bám yêu cầu thực hành soạn thảo văn bản. "
                    "Cô vẫn có thể chỉnh từng tiêu chí và điểm."
                )

                word_df = lay_barem_goi_y(
                    "Word",
                    mau_word
                )

                df_word = st.data_editor(
                    word_df,
                    num_rows="dynamic",
                    use_container_width=True,
                    column_config={
                        "loai":
                            st.column_config
                            .SelectboxColumn(
                                "Loại tiêu chí",
                                options=[
                                    "CO_NOI_DUNG",
                                    "SO_DOAN_TOI_THIEU",
                                    "SO_BANG_TOI_THIEU",
                                    "IN_DAM",
                                    "IN_NGHIENG",
                                    "FONT",
                                    "CO_CHU",
                                    "CAN_LE",
                                    "GIAN_DONG",
                                    "LE_TRANG",
                                    "SO_ANH_TOI_THIEU",
                                    "DAU_TRANG_CHAN_TRANG",
                                    "DANH_SACH_LIET_KE"
                                ]
                            ),

                        "diem":
                            st.column_config
                            .NumberColumn(
                                "Điểm",
                                min_value=0.0,
                                max_value=10.0,
                                step=0.25
                            ),

                        "gia_tri":
                            st.column_config
                            .TextColumn(
                                "Giá trị"
                            )
                    }
                )

                with st.expander(
                    "📖 Cách nhập barem Word"
                ):
                    st.markdown(
                        """
- `CO_NOI_DUNG`: các cụm bắt buộc, ngăn bởi `|`
- `SO_DOAN_TOI_THIEU`: ví dụ `5`
- `SO_BANG_TOI_THIEU`: ví dụ `1`
- `IN_DAM`: không cần giá trị
- `IN_NGHIENG`: không cần giá trị
- `FONT`: ví dụ `Times New Roman`
- `CO_CHU`: ví dụ `14`
- `CAN_LE`: `left`, `center`, `right`, `justify`
- `GIAN_DONG`: ví dụ `1.5`
- `LE_TRANG`: `trái,phải,trên,dưới`, ví dụ `3,2,2,2`\n- `SO_ANH_TOI_THIEU`: số ảnh tối thiểu\n- `DAU_TRANG_CHAN_TRANG`: kiểm tra có header/footer\n- `DANH_SACH_LIET_KE`: số đoạn danh sách tối thiểu
                        """
                    )

                if st.button(
                    "✅ Tạo bài Word và sinh mã",
                    type="primary"
                ):
                    try:
                        barem = (
                            tao_barem_dang_gia_tri(
                                ten_bai,
                                df_word
                            )
                        )

                        ma = tao_bai_moi(
                            ten_bai,
                            mo_ta,
                            "Word",
                            barem
                        )

                        st.success(
                            "Đã tạo bài!"
                        )
                        st.code(ma)

                    except Exception as e:
                        st.error(str(e))

            # ---------------------------------------------
            # POWERPOINT
            # ---------------------------------------------
            else:
                st.markdown(
                    "### Barem PowerPoint"
                )

                mau_ppt = st.selectbox(
                    "Barem chuẩn gợi ý",
                    list(BAREM_CHUAN_PPT.keys()),
                    key="mau_ppt"
                )

                st.caption(
                    "Barem gợi ý ưu tiên các yêu cầu Python có thể kiểm tra chắc chắn. "
                    "Các tiêu chí thẩm mỹ/phối màu cô có thể bổ sung khi chấm trực tiếp."
                )

                ppt_df = lay_barem_goi_y(
                    "PowerPoint",
                    mau_ppt
                )

                df_ppt = st.data_editor(
                    ppt_df,
                    num_rows="dynamic",
                    use_container_width=True,
                    column_config={
                        "loai":
                            st.column_config
                            .SelectboxColumn(
                                "Loại tiêu chí",
                                options=[
                                    "SO_SLIDE_TOI_THIEU",
                                    "TIEU_DE_SLIDE",
                                    "CO_NOI_DUNG",
                                    "SO_ANH_TOI_THIEU",
                                    "SO_BANG_TOI_THIEU",
                                    "FONT",
                                    "CO_CHU",
                                    "IN_DAM",
                                    "IN_NGHIENG",
                                    "CAN_LE"
                                ]
                            ),

                        "diem":
                            st.column_config
                            .NumberColumn(
                                "Điểm",
                                min_value=0.0,
                                max_value=10.0,
                                step=0.25
                            ),

                        "gia_tri":
                            st.column_config
                            .TextColumn(
                                "Giá trị"
                            )
                    }
                )

                with st.expander(
                    "📖 Cách nhập barem PowerPoint"
                ):
                    st.markdown(
                        """
- `SO_SLIDE_TOI_THIEU`: ví dụ `5`
- `TIEU_DE_SLIDE`: kiểm tra mỗi slide có tiêu đề; không cần giá trị
- `CO_NOI_DUNG`: các cụm bắt buộc, ngăn bởi `|`
- `SO_ANH_TOI_THIEU`: ví dụ `2`
- `SO_BANG_TOI_THIEU`: ví dụ `1`
- `FONT`: ví dụ `Arial`
- `CO_CHU`: ví dụ `24`
- `IN_DAM`: không cần giá trị
- `IN_NGHIENG`: không cần giá trị
- `CAN_LE`: `left`, `center`, `right`, `justify`
                        """
                    )

                if st.button(
                    "✅ Tạo bài PowerPoint và sinh mã",
                    type="primary"
                ):
                    try:
                        barem = (
                            tao_barem_dang_gia_tri(
                                ten_bai,
                                df_ppt
                            )
                        )

                        ma = tao_bai_moi(
                            ten_bai,
                            mo_ta,
                            "PowerPoint",
                            barem
                        )

                        st.success(
                            "Đã tạo bài!"
                        )
                        st.code(ma)

                    except Exception as e:
                        st.error(str(e))


        # =================================================
        # 17. QUẢN LÝ & KẾT QUẢ
        # =================================================

        with ql_tab:
            ds = lay_danh_sach_bai()

            if not ds:
                st.info(
                    "Chưa có bài."
                )

            else:
                opts = {
                    (
                        f"{x['code']} — "
                        f"{x['title']}"
                    ): x
                    for x in ds
                }

                key = st.selectbox(
                    "Chọn bài",
                    list(opts.keys())
                )

                bai = opts[key]

                loai = (
                    bai["rubric"]
                    .get(
                        "loai_file",
                        "Excel"
                    )
                )

                st.write(
                    f"**Mã:** "
                    f"`{bai['code']}`"
                )

                st.write(
                    f"**Loại:** {loai}"
                )

                st.write(
                    "**Trạng thái:** "
                    + (
                        "🟢 Đang mở"
                        if bai["active"]
                        else "🔴 Đã đóng"
                    )
                )

                if bai["active"]:
                    if st.button(
                        "🔒 Đóng nhận bài"
                    ):
                        cap_nhat_trang_thai_bai(
                            bai["code"],
                            False
                        )
                        st.rerun()
                else:
                    if st.button(
                        "🔓 Mở lại bài"
                    ):
                        cap_nhat_trang_thai_bai(
                            bai["code"],
                            True
                        )
                        st.rerun()

                ds_kq = lay_ket_qua_bai(
                    bai["code"]
                )

                if not ds_kq:
                    st.info(
                        "Chưa có học sinh "
                        "nộp bài."
                    )
                else:
                    bang = pd.DataFrame(
                        ds_kq
                    )

                    hien = (
                        bang[
                            [
                                "student_name",
                                "class_name",
                                "score",
                                "file_name",
                                "submitted_at"
                            ]
                        ]
                        .rename(
                            columns={
                                "student_name":
                                    "Họ và tên",

                                "class_name":
                                    "Lớp",

                                "score":
                                    "Điểm",

                                "file_name":
                                    "Tên file",

                                "submitted_at":
                                    "Thời gian nộp"
                            }
                        )
                    )

                    st.dataframe(
                        hien,
                        use_container_width=True,
                        hide_index=True
                    )

                    c1, c2, c3 = st.columns(3)

                    c1.metric(
                        "Số lượt nộp",
                        len(bang)
                    )

                    c2.metric(
                        "Điểm TB",
                        round(
                            float(
                                bang[
                                    "score"
                                ].mean()
                            ),
                            2
                        )
                    )

                    c3.metric(
                        "Điểm cao nhất",
                        float(
                            bang[
                                "score"
                            ].max()
                        )
                    )

                    file_excel = (
                        tao_file_excel_ket_qua(
                            ds_kq
                        )
                    )

                    if file_excel:
                        st.download_button(
                            "⬇️ Tải bảng điểm Excel",
                            data=file_excel,
                            file_name=(
                                f"Ket_qua_"
                                f"{bai['code']}.xlsx"
                            ),
                            mime=(
                                "application/"
                                "vnd.openxmlformats-"
                                "officedocument."
                                "spreadsheetml.sheet"
                            )
                        )

